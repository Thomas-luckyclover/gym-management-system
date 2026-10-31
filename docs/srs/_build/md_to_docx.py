"""Dựng docx theo định dạng SRS G5 từ file markdown đơn giản (tiêu đề, đoạn, danh sách, bảng, **đậm**, `code`).

Định dạng: Calibri; tiêu đề xanh 2E74B5; bảng viền đen, header hồng E6B8AF chữ đậm, chữ 11, giãn dòng 1.5.
Chạy: python3 md_to_docx.py input.md output.docx
"""
import re
import sys
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

FONT = "Calibri"
HEADER_FILL = "E6B8AF"
WIDTHS = {  # dxa, tổng 9360
    ("Field", "Value"): [2300, 7060],
    ("#", "API Name", "Calling Path", "Description"): [500, 2300, 1500, 5060],
    ("#", "Job Name", "Group", "Description"): [500, 2300, 1200, 5360],
    ("Scenario", "System behavior"): [2600, 6760],
    ("Job", "Schedule", "Use cases", "Main rules"): [3400, 1800, 2000, 2160],
    ("CBR-ID", "Rule Statement", "Use Case Relevant"): [1100, 6160, 2100],
    ("Date", "A*, M, D", "In charge", "Change Description"): [1300, 900, 1700, 5460],
    ("#", "Acronym", "Definition"): [500, 1500, 7360],
    ("#", "Chỗ", "Code hiện tại (develop)", "SRS"): [400, 1700, 4560, 2700],
}
HEAD = {1: (16, "2E74B5"), 2: (13, "2E74B5"), 3: (12, "1F4D78")}


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def font(run, size=12, bold=None, italic=None, color=None, mono=False):
    name = "Consolas" if mono else FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    f = rpr.find(qn("w:rFonts"))
    if f is None:
        f = el("w:rFonts")
        rpr.insert(0, f)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        f.set(qn(a), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(p, text, size=12, bold=False, italic=False):
    """**đậm**, `code`, \\| và \\* đã bỏ escape."""
    text = text.replace("\\|", "|")
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))", text):
        if not part:
            continue
        part = part.replace("\\*", "*") if not part.startswith("*") else part
        if part.startswith("**") and part.endswith("**"):
            font(p.add_run(part[2:-2]), size, bold=True, italic=italic or None)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            font(p.add_run(part[1:-1]), size, bold=bold or None, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            font(p.add_run(part[1:-1]), size - 1, bold=bold or None, mono=True)
        else:
            font(p.add_run(part), size, bold=bold or None, italic=italic or None)


def para(doc, text="", size=12, bold=False, italic=False, indent=0, after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15
    if indent:
        pf.left_indent = Cm(indent)
    if text:
        add_inline(p, text, size, bold, italic)
    return p


def heading(doc, level, text):
    size, color = HEAD.get(level, (12, "1F4D78"))
    p = doc.add_paragraph()
    p.style = doc.styles["Heading %d" % level]
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 10)
    pf.space_after = Pt(6)
    font(p.add_run(text), size, bold=True, color=color)


def list_item(doc, marker, text, level=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.9 + 0.8 * level)
    pf.first_line_indent = Cm(-0.6)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    font(p.add_run(marker + "\t"), 12)
    pf.tab_stops.add_tab_stop(Cm(0.9 + 0.8 * level))
    add_inline(p, text)


def table(doc, rows):
    header = tuple(c.strip().replace("\\*", "*").replace("\\|", "|") for c in rows[0])
    n = len(header)
    widths = WIDTHS.get(header) or [9360 // n] * n
    t = doc.add_table(rows=len(rows), cols=n)
    tbl = t._tbl
    tblPr = tbl.tblPr
    tblPr.append(el("w:tblW", **{"w:type": "dxa", "w:w": "9360"}))
    borders = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(el("w:" + side, **{"w:val": "single", "w:sz": "4", "w:color": "000000"}))
    tblPr.append(borders)
    tblPr.append(el("w:tblLayout", **{"w:type": "fixed"}))
    for i, gc in enumerate(tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(widths[i]))
    for ri, row in enumerate(rows):
        tr = t.rows[ri]
        if ri == 0:
            tr._tr.get_or_add_trPr().append(el("w:tblHeader"))
        for ci in range(n):
            cell = tr.cells[ci]
            tcPr = cell._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcW")):
                tcPr.remove(old)
            tcPr.insert(0, el("w:tcW", **{"w:type": "dxa", "w:w": str(widths[ci])}))
            if ri == 0:
                tcPr.append(el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": HEADER_FILL}))
            mar = el("w:tcMar")
            for side, w in (("top", "40"), ("left", "100"), ("bottom", "40"), ("right", "100")):
                mar.append(el("w:" + side, **{"w:type": "dxa", "w:w": w}))
            tcPr.append(mar)
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.5
            text = row[ci].strip() if ci < len(row) else ""
            add_inline(p, text, 11, bold=(ri == 0))
    para(doc, after=4)


def split_row(line):
    line = line.strip().strip("|")
    cells, cur, i = [], "", 0
    while i < len(line):
        if line[i] == "\\" and i + 1 < len(line):
            cur += line[i:i + 2]
            i += 2
            continue
        if line[i] == "|":
            cells.append(cur)
            cur = ""
        else:
            cur += line[i]
        i += 1
    cells.append(cur)
    return cells


def main(src, dst):
    lines = open(src, encoding="utf-8").read().splitlines()
    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1)
        title = next((l.split(":", 1)[1].strip().strip('"') for l in lines[1:end] if l.startswith("title:")), "")
        lines = lines[end + 1:]
    else:
        title = ""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.54)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    if title:
        p = doc.add_paragraph()
        font(p.add_run(title), 18, bold=True, color="1F4D78")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        m = re.match(r"^(#{1,3}) (.*)", s)
        if m:
            heading(doc, len(m.group(1)), m.group(2))
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [split_row(b) for b in block if not re.match(r"^\|\s*:?-{2,}", b.strip())]
            table(doc, rows)
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.|[a-z]\.)\s+(.*)", line)
        if m:
            level = 1 if len(m.group(1)) >= 2 else 0
            marker = "•" if m.group(2) in "-*" else m.group(2)
            list_item(doc, marker, m.group(3), level)
            i += 1
            continue
        if s.startswith(">"):
            para(doc, s.lstrip("> ").strip(), italic=False, indent=0.8)
            i += 1
            continue
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^\s*(#|\||[-*] |\d+\. |>)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        para(doc, " ".join(buf).replace("\\.", "."))
    doc.save(dst)
    print("saved", dst, "tables:", len(doc.tables))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
